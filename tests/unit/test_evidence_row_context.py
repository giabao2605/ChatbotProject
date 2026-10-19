import pytest
from sqlalchemy import text

from mech_chatbot.adapters.protected_files import FilesystemProtectedFileStorage
from mech_chatbot.db.repositories.structured_evidence import save_extracted_scope
from mech_chatbot.ingestion.pdf.structured import extract_structured_table_file
from mech_chatbot.services.rag_evidence_service import list_scope_facts
from .test_structured_evidence_repository import PROFILE, store


@pytest.fixture
def source(store, tmp_path):
    raw = tmp_path / 'raw'
    raw.mkdir()
    path = raw / 'ledger.csv'
    path.write_text('ID,Date,Amount,Status\n001,2026-10-01,12.50,posted\n001,2026-10-02,-2.50,cancelled\n', encoding='utf-8')
    scope = extract_structured_table_file(path, doc_id='1', version='2', snapshot_id='row-context').tables[0].scopes[0]
    save_extracted_scope(scope, db_engine=store)
    with store.begin() as conn:
        conn.execute(text('UPDATE TaiLieu SET FilePath=:path'), {'path': str(path)})
    storage = FilesystemProtectedFileStorage(project_root=tmp_path, raw_root=raw, processed_root=tmp_path / 'processed')
    return scope, path, storage


def test_original_row_context_is_paged_and_preserves_duplicate_business_ids(store, source):
    scope, path, storage = source
    first = list_scope_facts(scope.scope_id, PROFILE, offset=0, limit=1, db_engine=store, source_storage=storage)
    second = list_scope_facts(scope.scope_id, PROFILE, offset=1, limit=1, db_engine=store, source_storage=storage)
    assert first['has_more'] and not second['has_more']
    assert first['items'][0]['source_row'] == {'headers': ['ID', 'Date', 'Amount', 'Status'],
        'cells': ['001', '2026-10-01', '12.50', 'posted'], 'source_row': 2}
    assert second['items'][0]['source_row']['cells'] == ['001', '2026-10-02', '-2.50', 'cancelled']
    assert str(path) not in str(first)
    assert first['items'][0]['review_status'] == 'pending'


@pytest.mark.parametrize('fault', ['changed', 'missing', 'outside'])
def test_unavailable_file_never_supplies_context(store, source, tmp_path, fault):
    scope, path, storage = source
    if fault == 'changed':
        path.write_text('ID,Amount\nsecret,999\n')
    elif fault == 'missing':
        path.unlink()
    else:
        other = tmp_path / 'outside.csv'
        other.write_bytes(path.read_bytes())
        with store.begin() as conn:
            conn.execute(text('UPDATE TaiLieu SET FilePath=:path'), {'path': str(other)})
    page = list_scope_facts(scope.scope_id, PROFILE, db_engine=store, source_storage=storage)
    assert all(item['source_row'] is None for item in page['items'])
    assert page['items'][0]['raw_value'] == '12.50'


def test_review_acl_is_checked_before_touching_source_file(store, source, monkeypatch):
    scope, _, storage = source
    monkeypatch.setattr(storage, 'resolve_original', lambda *_: pytest.fail('unauthorized file access'))
    with pytest.raises(PermissionError):
        list_scope_facts(scope.scope_id, {**PROFILE, 'allowed_sites': ['other']}, db_engine=store, source_storage=storage)


def test_pending_document_review_keeps_access_to_its_original_row(store, source):
    scope, _, storage = source
    with store.begin() as conn:
        conn.execute(text("UPDATE TaiLieu SET LifecycleStatus='pending_review',ReviewStatus='pending',PublicationState='draft',Servable=0"))
    page = list_scope_facts(scope.scope_id, PROFILE, db_engine=store, source_storage=storage)
    assert page['items'][0]['source_row']['cells'][0] == '001'


def test_parser_reads_the_hashed_copy_if_original_is_replaced(store, source, monkeypatch):
    from mech_chatbot.services import evidence_row_context as context
    scope, path, storage = source
    original = context.extract_structured_table_file
    def replace_source_then_parse(copy, **kwargs):
        path.write_text('ID,Amount\nchanged,999\n')
        assert copy != path
        return original(copy, **kwargs)
    monkeypatch.setattr(context, 'extract_structured_table_file', replace_source_then_parse)
    page = list_scope_facts(scope.scope_id, PROFILE, db_engine=store, source_storage=storage)
    assert page['items'][0]['source_row']['cells'][0] == '001'


@pytest.mark.parametrize('fault', ['oversize', 'parse_error'])
def test_bounded_source_failures_preserve_existing_facts(store, source, monkeypatch, fault):
    from mech_chatbot.services import evidence_row_context as context
    scope, _, storage = source
    if fault == 'oversize':
        monkeypatch.setattr(context, 'MAX_UPLOAD_BYTES', 5)
    else:
        def invalid(*args, **kwargs):
            raise ValueError('private file details')
        monkeypatch.setattr(context, 'extract_structured_table_file', invalid)
    page = list_scope_facts(scope.scope_id, PROFILE, db_engine=store, source_storage=storage)
    assert page['items'][0]['source_row'] is None
    assert 'private file details' not in str(page)


def test_xlsx_context_keeps_sheet_identity_and_literal_text(store, tmp_path):
    from openpyxl import Workbook
    book = Workbook()
    book.active.title = 'First'
    book.create_sheet('Second')
    for sheet, amount in zip(book.worksheets, (10, 20)):
        sheet.append(['ID', 'Amount', 'Status'])
        sheet.append(['001', amount, sheet.title])
    path = tmp_path / 'ledger.xlsx'
    book.save(path)
    document = extract_structured_table_file(path, doc_id='1', version='2', snapshot_id='xlsx-context')
    for table in document.tables:
        save_extracted_scope(table.scopes[0], db_engine=store)
    with store.begin() as conn:
        conn.execute(text('UPDATE TaiLieu SET FilePath=:path'), {'path': str(path)})
    storage = FilesystemProtectedFileStorage(project_root=tmp_path, raw_root=tmp_path, processed_root=tmp_path)
    page = list_scope_facts(document.tables[1].scopes[0].scope_id, PROFILE, db_engine=store, source_storage=storage)
    assert page['items'][0]['source_row']['cells'] == ['001', '20', 'Second']
